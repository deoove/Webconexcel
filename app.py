from flask import Flask, jsonify, request, send_from_directory
import openpyxl
import os

app = Flask(__name__, static_folder="static")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "datos.xlsx")


def get_workbook():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Personas"
        ws.append(["ID", "Nombre", "Apellido", "DNI"])
        wb.save(EXCEL_FILE)
    return openpyxl.load_workbook(EXCEL_FILE)


def save_workbook(wb):
    wb.save(EXCEL_FILE)


def get_next_id(ws):
    ids = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]
    return max(ids, default=0) + 1


def get_all_personas():
    wb = get_workbook()
    ws = wb.active
    personas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            personas.append({"id": row[0], "nombre": row[1], "apellido": row[2], "dni": str(row[3])})
    return personas


def find_row(ws, persona_id):
    for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if row[0].value == persona_id:
            return idx
    return None


def dni_existe(dni, excluir_id=None):
    for p in get_all_personas():
        if str(p["dni"]) == dni and p["id"] != excluir_id:
            return True
    return False


def validar(nombre, apellido, dni, excluir_id=None):
    if not nombre or not apellido or not dni:
        return "Todos los campos son obligatorios."
    if not dni.isdigit():
        return "El DNI debe contener solo números."
    if dni_existe(dni, excluir_id):
        return "Ya existe una persona con ese DNI."
    return None


@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.get("/api/personas")
def listar():
    return jsonify(get_all_personas())


@app.post("/api/personas")
def crear():
    data = request.get_json(silent=True) or {}
    nombre = str(data.get("nombre", "")).strip()
    apellido = str(data.get("apellido", "")).strip()
    dni = str(data.get("dni", "")).strip()
    error = validar(nombre, apellido, dni)
    if error:
        return jsonify({"error": error}), 400
    wb = get_workbook()
    ws = wb.active
    new_id = get_next_id(ws)
    ws.append([new_id, nombre, apellido, dni])
    save_workbook(wb)
    return jsonify({"id": new_id, "nombre": nombre, "apellido": apellido, "dni": dni}), 201


@app.put("/api/personas/<int:persona_id>")
def actualizar(persona_id):
    wb = get_workbook()
    ws = wb.active
    row_idx = find_row(ws, persona_id)
    if row_idx is None:
        return jsonify({"error": "Persona no encontrada."}), 404
    data = request.get_json(silent=True) or {}
    nombre = str(data.get("nombre", "")).strip()
    apellido = str(data.get("apellido", "")).strip()
    dni = str(data.get("dni", "")).strip()
    error = validar(nombre, apellido, dni, excluir_id=persona_id)
    if error:
        return jsonify({"error": error}), 400
    ws.cell(row=row_idx, column=2).value = nombre
    ws.cell(row=row_idx, column=3).value = apellido
    ws.cell(row=row_idx, column=4).value = dni
    save_workbook(wb)
    return jsonify({"id": persona_id, "nombre": nombre, "apellido": apellido, "dni": dni})


@app.delete("/api/personas/<int:persona_id>")
def eliminar(persona_id):
    wb = get_workbook()
    ws = wb.active
    row_idx = find_row(ws, persona_id)
    if row_idx is None:
        return jsonify({"error": "Persona no encontrada."}), 404
    ws.delete_rows(row_idx)
    save_workbook(wb)
    return jsonify({"ok": True})


if __name__ == "__main__":
    get_workbook()  # crea datos.xlsx si todavía no existe
    print(f"Base de datos Excel: {EXCEL_FILE}")
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=True)
